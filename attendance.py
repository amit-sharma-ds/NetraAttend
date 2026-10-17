"""Attendance rules (on time / late / half day), marking and reports."""
import io
import re
import threading
import time
from collections import Counter, deque
from datetime import date, datetime, timedelta

from pymongo.errors import DuplicateKeyError, PyMongoError

import config
from db import db, get_settings, now_str, put_photo

VISIT_GAP_MIN = 3     # not seen for this long -> the next sighting is a new movement (OUT)
MAX_PENDING = 3000    # sightings kept in memory while the database is unreachable
EXIT_AFTER_IN_SEC = 120   # exit camera ignores a person this soon after their check-in (turning at the door)


def parse_hm(value):
    return datetime.strptime(value.strip(), "%H:%M").time()


def parse_clock(value):
    value = value.strip()
    return datetime.strptime(value, "%H:%M:%S" if value.count(":") == 2 else "%H:%M").time()


def _minutes_between(a, b):
    """b - a in minutes for two datetime.time values."""
    base = date(2000, 1, 1)
    return (datetime.combine(base, b) - datetime.combine(base, a)).total_seconds() / 60


class Rules:
    def __init__(self, settings=None):
        s = settings or get_settings()
        self.office_start = parse_hm(s["office_start"])
        self.grace = int(float(s["grace_min"]))
        self.half_day_after = parse_hm(s["half_day_after"])
        self.half_day_hours = float(s["half_day_hours"])
        self.min_out_gap = int(float(s["min_out_gap_min"]))
        self.weekly_off = {int(x) for x in s["weekly_off"].replace(" ", "").split(",") if x.isdigit()}
        self.holidays = set(re.findall(r"\d{4}-\d{2}-\d{2}", s.get("holidays", "")))

    @property
    def late_after(self):
        return (datetime.combine(date(2000, 1, 1), self.office_start)
                + timedelta(minutes=self.grace)).time()

    def day_off(self, day):
        if day.isoformat() in self.holidays:
            return "Holiday"
        if day.weekday() in self.weekly_off:
            return "Weekly Off"
        return None

    def evaluate(self, day, check_in, check_out, today):
        cin = parse_clock(check_in)
        late_min = max(0, int(_minutes_between(self.office_start, cin)))
        hours = None
        if check_out:
            worked = _minutes_between(cin, parse_clock(check_out)) / 60
            hours = round(worked, 2) if worked > 0 else None
        if cin > self.half_day_after:
            status = "Half Day"
        elif late_min > self.grace:
            status = "Late"
        else:
            status = "On Time"
        # hours rule only once the day is over (people are still at work today)
        if day < today and hours is not None and hours < self.half_day_hours:
            status = "Half Day"
        # late time is only counted on "Late" days; a Half Day is its own category, so a
        # 17:00 arrival doesn't add hours of "late" to the totals
        return {"status": status, "late_min": late_min if status == "Late" else 0, "hours": hours}


class Recorder:
    """Turns a recognised face into attendance (thread-safe).

    - First sighting of the day             -> IN (check-in time + photo)
    - Each new appearance later in the day  -> OUT movement, logged with time + photo;
      check-out = latest movement / last seen. Hours = first IN to last OUT.
    A "new appearance" means the person was not seen for VISIT_GAP_MIN minutes,
    so standing in front of the camera for a while is one movement, not many.

    If MongoDB is unreachable (internet down) sightings wait in memory, with
    their original time, and are written by flush_pending() once it is back.
    """

    def __init__(self):
        self.lock = threading.Lock()
        self.last = {}            # emp_id -> (monotonic time, result)
        self.last_seen = {}       # emp_id -> datetime of last sighting
        self.pending = deque(maxlen=MAX_PENDING)
        self.pending_lock = threading.Lock()
        self._rules = None
        self._rules_at = 0.0

    def invalidate_rules(self):
        self._rules = None

    def rules(self):
        if self._rules is None or time.monotonic() - self._rules_at > 30:
            self._rules = Rules()
            self._rules_at = time.monotonic()
        return self._rules

    @staticmethod
    def _save_snapshot(snap_jpg, day, emp_id, kind):
        if not snap_jpg:
            return None
        return put_photo(f"{day}/{emp_id}_{kind}.jpg", snap_jpg, emp_id)

    @staticmethod
    def _seed_last_seen(emp_id, day):
        """After a restart, continue from what the database knows about today."""
        row = db().attendance.find_one({"emp_id": emp_id, "date": day.isoformat()},
                                       {"check_in": 1, "check_out": 1})
        if row is None:
            return None
        return datetime.combine(day, parse_clock(row.get("check_out") or row["check_in"]))

    def record(self, emp_id, when, score, snap_jpg, direction=None):
        """direction None = the only camera (IN and OUT as above); "in" / "out" = the
        entrance / exit camera when an exit camera is set (see _record_dir)."""
        try:
            return self._record(emp_id, when, score, snap_jpg, direction)
        except PyMongoError as exc:
            with self.pending_lock:
                if not self.pending:
                    print(f"Database unreachable ({type(exc).__name__}) - holding attendance "
                          f"in memory until it is back", flush=True)
                self.pending.append((emp_id, when, score, snap_jpg, direction))
            return {"kind": "QUEUED", "time": when.strftime("%H:%M"), "status": "",
                    "late_min": 0, "photo": None, "fresh": False}

    def flush_pending(self):
        """Write sightings held while the database was down, oldest first."""
        while True:
            with self.pending_lock:
                if not self.pending:
                    return
                item = self.pending[0]
            try:
                self._record(*item, replay=True)
            except PyMongoError:
                return
            with self.pending_lock:
                if self.pending and self.pending[0] is item:
                    self.pending.popleft()
                if not self.pending:
                    print("Database back - queued attendance saved", flush=True)

    def _record(self, emp_id, when, score, snap_jpg, direction=None, replay=False):
        if direction in ("in", "out"):
            return self._record_dir(emp_id, when, score, snap_jpg, direction, replay)
        with self.lock:
            d = db()
            day = when.strftime("%Y-%m-%d")
            clock = when.strftime("%H:%M:%S")
            last = self.last_seen.get(emp_id)
            if last is None or last.date() != when.date():
                last = self._seed_last_seen(emp_id, when.date())
            new_visit = last is None or when - last >= timedelta(minutes=VISIT_GAP_MIN)

            now_mono = time.monotonic()
            prev = self.last.get(emp_id)
            if not replay and not new_visit and prev and now_mono - prev[0] < config.COOLDOWN_SEC:
                self.last_seen[emp_id] = when
                return dict(prev[1], fresh=False)

            rules = self.rules()
            row = d.attendance.find_one({"emp_id": emp_id, "date": day})
            fresh = new_visit
            if row is None:
                photo = self._save_snapshot(snap_jpg, day, emp_id, "in")
                try:
                    d.attendance.insert_one({"emp_id": emp_id, "date": day, "check_in": clock,
                                             "check_out": None, "in_photo": photo, "out_photo": None,
                                             "in_score": score, "source": "camera", "note": ""})
                except DuplicateKeyError:   # written by an earlier attempt that timed out
                    pass
                d.events.insert_one({"emp_id": emp_id, "ts": f"{day} {clock}", "kind": "IN",
                                     "score": score, "photo": photo})
                ev = rules.evaluate(when.date(), clock, None, when.date())
                result = {"kind": "IN", "time": clock[:5], "status": ev["status"],
                          "late_min": ev["late_min"], "photo": photo}
                fresh = True
            else:
                check_in = datetime.combine(when.date(), parse_clock(row["check_in"]))
                if when - check_in >= timedelta(minutes=rules.min_out_gap):
                    if new_visit:
                        # a new movement: own photo + log entry
                        photo = self._save_snapshot(snap_jpg, day, emp_id, when.strftime("%H%M%S"))
                        d.attendance.update_one({"_id": row["_id"]},
                                                {"$set": {"check_out": clock, "out_photo": photo}})
                        d.events.insert_one({"emp_id": emp_id, "ts": f"{day} {clock}", "kind": "OUT",
                                             "score": score, "photo": photo})
                    else:
                        # same movement, still in view: keep last-seen time current
                        photo = None
                        d.attendance.update_one({"_id": row["_id"]}, {"$set": {"check_out": clock}})
                    hours = round((when - check_in).total_seconds() / 3600, 2)
                    result = {"kind": "OUT", "time": clock[:5], "status": "", "late_min": 0,
                              "photo": photo, "hours": hours}
                else:
                    ev = rules.evaluate(when.date(), row["check_in"], None, when.date())
                    result = {"kind": "SEEN", "time": row["check_in"][:5],
                              "status": ev["status"], "late_min": ev["late_min"], "photo": None}
            # only remember the sighting once it is safely in the database
            self.last_seen[emp_id] = when
            self.last[emp_id] = (now_mono, result)
            return dict(result, fresh=fresh)

    @staticmethod
    def _seed_last_event(emp_id, day, kind):
        """After a restart: time of today's last camera IN / OUT movement."""
        ev = db().events.find_one({"emp_id": emp_id, "kind": kind, "source": {"$exists": False},
                                   "ts": {"$gte": f"{day} 00:00:00", "$lte": f"{day} 23:59:59"}},
                                  sort=[("ts", -1)])
        return datetime.strptime(ev["ts"], "%Y-%m-%d %H:%M:%S") if ev else None

    def _record_dir(self, emp_id, when, score, snap_jpg, direction, replay):
        """Two cameras: the entrance camera marks IN, the exit camera marks OUT.

        - entrance camera: first sighting of the day = check-in. Coming back in after
          an OUT = a new IN movement (photo + log). Seen again while already in = nothing.
        - exit camera: every new appearance = an OUT movement (photo + log) and the
          latest check-out; standing in view keeps the check-out time current.
          Hours stay first IN to last OUT, same as with one camera.
        - the two stay separate: the exit camera never marks IN. Someone it sees with no
          IN today is only logged as OUT (photo + activity) so the admin can fix the day.
        """
        with self.lock:
            d = db()
            day = when.strftime("%Y-%m-%d")
            clock = when.strftime("%H:%M:%S")
            key = (emp_id, direction)
            last = self.last_seen.get(key)
            if last is None or last.date() != when.date():
                last = self._seed_last_event(emp_id, when.date(), "IN" if direction == "in" else "OUT")
            new_visit = last is None or when - last >= timedelta(minutes=VISIT_GAP_MIN)

            now_mono = time.monotonic()
            prev = self.last.get(key)
            if not replay and not new_visit and prev and now_mono - prev[0] < config.COOLDOWN_SEC:
                self.last_seen[key] = when
                return dict(prev[1], fresh=False)

            rules = self.rules()
            row = d.attendance.find_one({"emp_id": emp_id, "date": day})
            fresh = new_visit
            if row is None and direction == "out":
                photo = None
                if new_visit:
                    photo = self._save_snapshot(snap_jpg, day, emp_id, when.strftime("%H%M%S"))
                    d.events.insert_one({"emp_id": emp_id, "ts": f"{day} {clock}", "kind": "OUT",
                                         "score": score, "photo": photo})
                result = {"kind": "OUT", "time": clock[:5], "status": "No IN today", "late_min": 0,
                          "photo": photo, "hours": "-"}
            elif row is None:
                photo = self._save_snapshot(snap_jpg, day, emp_id, "in")
                try:
                    d.attendance.insert_one({"emp_id": emp_id, "date": day, "check_in": clock,
                                             "check_out": None, "in_photo": photo, "out_photo": None,
                                             "in_score": score, "source": "camera", "note": ""})
                except DuplicateKeyError:
                    pass
                d.events.insert_one({"emp_id": emp_id, "ts": f"{day} {clock}", "kind": "IN",
                                     "score": score, "photo": photo})
                ev = rules.evaluate(when.date(), clock, None, when.date())
                result = {"kind": "IN", "time": clock[:5], "status": ev["status"],
                          "late_min": ev["late_min"], "photo": photo}
                fresh = True
            elif direction == "in":
                latest = d.events.find_one({"emp_id": emp_id, "ts": {"$gte": f"{day} 00:00:00",
                                                                      "$lte": f"{day} {clock}"}},
                                           sort=[("ts", -1)])
                if new_visit and latest is not None and latest["kind"] == "OUT":
                    photo = self._save_snapshot(snap_jpg, day, emp_id, when.strftime("%H%M%S"))
                    d.events.insert_one({"emp_id": emp_id, "ts": f"{day} {clock}", "kind": "IN",
                                         "score": score, "photo": photo})
                    result = {"kind": "IN", "time": clock[:5], "status": "Back in", "late_min": 0,
                              "photo": photo}
                else:
                    ev = rules.evaluate(when.date(), row["check_in"], None, when.date())
                    result = {"kind": "SEEN", "time": row["check_in"][:5],
                              "status": ev["status"], "late_min": ev["late_min"], "photo": None}
            else:
                check_in = datetime.combine(when.date(), parse_clock(row["check_in"]))
                if when - check_in < timedelta(seconds=EXIT_AFTER_IN_SEC):
                    ev = rules.evaluate(when.date(), row["check_in"], None, when.date())
                    result = {"kind": "SEEN", "time": row["check_in"][:5],
                              "status": ev["status"], "late_min": ev["late_min"], "photo": None}
                    fresh = False
                else:
                    if new_visit:
                        photo = self._save_snapshot(snap_jpg, day, emp_id, when.strftime("%H%M%S"))
                        d.attendance.update_one({"_id": row["_id"]},
                                                {"$set": {"check_out": clock, "out_photo": photo}})
                        d.events.insert_one({"emp_id": emp_id, "ts": f"{day} {clock}", "kind": "OUT",
                                             "score": score, "photo": photo})
                    else:
                        photo = None
                        d.attendance.update_one({"_id": row["_id"]}, {"$set": {"check_out": clock}})
                    hours = round((when - check_in).total_seconds() / 3600, 2)
                    result = {"kind": "OUT", "time": clock[:5], "status": "", "late_min": 0,
                              "photo": photo, "hours": hours}
            self.last_seen[key] = when
            self.last[key] = (now_mono, result)
            return dict(result, fresh=fresh)


def manual_entry(emp_id, day, check_in, check_out, note):
    """Admin correction for a day. Empty check_in deletes the record."""
    attendance = db().attendance
    if not check_in:
        attendance.delete_one({"emp_id": emp_id, "date": day})
        return
    cin = parse_clock(check_in).strftime("%H:%M:%S")
    cout = parse_clock(check_out).strftime("%H:%M:%S") if check_out else None
    if cout and cout <= cin:
        raise ValueError("Check-out must be after check-in.")
    attendance.update_one(
        {"emp_id": emp_id, "date": day},
        {"$set": {"check_in": cin, "check_out": cout, "source": "manual", "note": note},
         "$setOnInsert": {"in_photo": None, "out_photo": None, "in_score": None}},
        upsert=True)


def _daterange(start, end):
    d = start
    while d <= end:
        yield d
        d += timedelta(days=1)


def _fmt_clock(value):
    return value[:5] if value else ""


def employees(query=None):
    """Employee documents sorted by name, with "id" = _id for templates."""
    return [dict(e, id=e["_id"]) for e in
            db().employees.find(query or {}, {"password_hash": 0}).sort("name", 1)]


def _movements(start, end):
    """{(emp_id, 'YYYY-MM-DD'): [{"time", "kind", "photo"}, ...]} in time order."""
    out = {}
    for r in db().events.find({"ts": {"$gte": f"{start} 00:00:00", "$lte": f"{end} 23:59:59"}},
                              {"_id": 0}).sort("ts", 1):
        out.setdefault((r["emp_id"], r["ts"][:10]), []).append(
            {"time": r["ts"][11:16], "kind": r["kind"], "photo": r.get("photo"),
             "field": r.get("source") == "field", "lat": r.get("lat"), "lng": r.get("lng"),
             "verified": r.get("verified"), "device": r.get("device") if r.get("source") == "device" else None})
    return out


def moves_text(moves):
    parts = []
    for m in moves:
        text = f"{m['kind']} {m['time']}"
        if m.get("field"):
            text += f" (field {m['lat']:.5f},{m['lng']:.5f})" if m.get("lat") is not None else " (field)"
        parts.append(text)
    return ", ".join(parts)


def record_field(emp_id, when, kind, photo, lat, lng, accuracy, note, verified, score):
    """Attendance marked by the employee from outside the office (selfie + GPS).
    Same day record as the camera: first IN = check-in, every OUT = latest check-out."""
    d = db()
    day = when.strftime("%Y-%m-%d")
    clock = when.strftime("%H:%M:%S")
    row = d.attendance.find_one({"emp_id": emp_id, "date": day})
    last = d.events.find_one({"emp_id": emp_id, "ts": {"$gte": f"{day} 00:00:00"}}, sort=[("ts", -1)])
    if last and last["kind"] == kind and when - datetime.strptime(last["ts"], "%Y-%m-%d %H:%M:%S") \
            < timedelta(minutes=2):
        raise ValueError(f"Already marked {kind} at {last['ts'][11:16]}.")
    if kind == "OUT" and row is None:
        raise ValueError("Mark IN first.")
    if row is None:
        try:
            d.attendance.insert_one({"emp_id": emp_id, "date": day, "check_in": clock, "check_out": None,
                                     "in_photo": photo, "out_photo": None, "in_score": score,
                                     "source": "field", "note": note})
        except DuplicateKeyError:
            pass
    elif kind == "OUT":
        d.attendance.update_one({"_id": row["_id"]}, {"$set": {"check_out": clock, "out_photo": photo}})
    d.events.insert_one({"emp_id": emp_id, "ts": f"{day} {clock}", "kind": kind, "score": score,
                         "photo": photo, "source": "field", "lat": lat, "lng": lng,
                         "accuracy": accuracy, "note": note, "verified": verified})
    return clock[:5]


DOUBLE_PUNCH_SEC = 120   # a second machine punch this soon after the first is the same punch


def record_punch(emp_id, when, device, kind=None):
    """Punch from an external machine / system (API). Machines often resend old logs and
    in any order, so the day is rebuilt as earliest punch = IN, latest punch = OUT.
    Returns {"result": "saved" | "duplicate", "kind": "IN" | "OUT"}."""
    d = db()
    day = when.strftime("%Y-%m-%d")
    clock = when.strftime("%H:%M:%S")
    ts = f"{day} {clock}"
    if d.events.find_one({"emp_id": emp_id, "ts": ts, "source": "device", "device": device}, {"_id": 1}):
        return {"result": "duplicate", "kind": None}
    row = d.attendance.find_one({"emp_id": emp_id, "date": day})
    if row is None:
        kind = kind or "IN"
        try:
            d.attendance.insert_one({"emp_id": emp_id, "date": day, "check_in": clock, "check_out": None,
                                     "in_photo": None, "out_photo": None, "in_score": None,
                                     "source": "device", "note": ""})
        except DuplicateKeyError:
            row = d.attendance.find_one({"emp_id": emp_id, "date": day})
    if row is not None:
        first = datetime.combine(when.date(), parse_clock(row["check_in"]))
        last = datetime.combine(when.date(), parse_clock(row["check_out"])) if row.get("check_out") else first
        if abs((when - first).total_seconds()) < DOUBLE_PUNCH_SEC and when >= first:
            kind = kind or "IN"           # double punch right after IN: keep it as IN, change nothing
        elif when < first:
            kind = kind or "IN"           # an older punch arrived late: it is the real IN
            d.attendance.update_one({"_id": row["_id"]}, {"$set": {
                "check_in": clock, "check_out": row.get("check_out") or row["check_in"]}})
        else:
            kind = kind or "OUT"
            if when >= last:
                d.attendance.update_one({"_id": row["_id"]}, {"$set": {"check_out": clock}})
    d.events.insert_one({"emp_id": emp_id, "ts": ts, "kind": kind, "score": None, "photo": None,
                         "source": "device", "device": device})
    return {"result": "saved", "kind": kind}


# ------------------------------------------------------------------ leave
def parse_date(value):
    return datetime.strptime(value, "%Y-%m-%d").date()


def leave_types(settings=None):
    """[(name, days per year)] from Settings; 0 days = no limit (e.g. unpaid)."""
    out = []
    for part in re.split(r"[,\n]", (settings or get_settings()).get("leave_types", "")):
        name, _, quota = part.partition(":")
        if name.strip():
            try:
                out.append((name.strip(), float(quota) if quota.strip() else 0.0))
            except ValueError:
                out.append((name.strip(), 0.0))
    return out or [("Casual Leave", 12.0), ("Sick Leave", 6.0), ("Unpaid Leave", 0.0)]


def approved_leaves(start, end, emp_id=None):
    """{(emp_id, 'YYYY-MM-DD'): {"type", "half"}} for approved leave between start and end."""
    query = {"status": "approved", "from": {"$lte": end.isoformat()}, "to": {"$gte": start.isoformat()}}
    if emp_id:
        query["emp_id"] = emp_id
    out = {}
    for lv in db().leaves.find(query):
        for d in _daterange(max(start, parse_date(lv["from"])), min(end, parse_date(lv["to"]))):
            out[(lv["emp_id"], d.isoformat())] = {"type": lv["type"], "half": bool(lv.get("half_day"))}
    return out


def leave_balance(emp_id, year):
    used, pending = Counter(), Counter()
    for lv in db().leaves.find({"emp_id": emp_id, "status": {"$in": ["approved", "pending"]},
                                "from": {"$gte": f"{year}-01-01", "$lte": f"{year}-12-31"}}):
        (used if lv["status"] == "approved" else pending)[lv["type"]] += lv.get("days", 0)
    return [{"type": t, "quota": q, "used": used[t], "pending": pending[t],
             "left": max(0.0, q - used[t]) if q else None} for t, q in leave_types()]


def apply_leave(emp_id, ltype, start, end, half_day, reason):
    """Employee leave request -> pending. Returns the number of leave days."""
    types = dict(leave_types())
    today = date.today()
    if ltype not in types:
        raise ValueError("Choose a leave type.")
    if end < start:
        raise ValueError("'To' date is before 'From' date.")
    if start < today - timedelta(days=30) or end > today + timedelta(days=180):
        raise ValueError("Leave can be applied from 30 days back up to 6 months ahead.")
    if start.year != end.year:
        raise ValueError("Leave crossing into the next year: apply as two requests.")
    if half_day and start != end:
        raise ValueError("Half-day leave must be for one date.")
    if not reason.strip():
        raise ValueError("Write a reason.")
    rules = Rules()
    days = 0.5 if half_day else float(sum(1 for d in _daterange(start, end) if not rules.day_off(d)))
    if not days:
        raise ValueError("Those dates are weekly off / holidays - no leave needed.")
    clash = db().leaves.find_one({"emp_id": emp_id, "status": {"$in": ["pending", "approved"]},
                                  "from": {"$lte": end.isoformat()}, "to": {"$gte": start.isoformat()}})
    if clash:
        raise ValueError(f"You already have a {clash['status']} leave from {clash['from']} to {clash['to']}.")
    quota = types[ltype]
    if quota:
        bal = next(b for b in leave_balance(emp_id, start.year) if b["type"] == ltype)
        left = quota - bal["used"] - bal["pending"]
        if days > left:
            raise ValueError(f"Only {left:g} day(s) of {ltype} left this year (pending requests included).")
    db().leaves.insert_one({"emp_id": emp_id, "type": ltype, "from": start.isoformat(), "to": end.isoformat(),
                            "half_day": bool(half_day), "days": days, "reason": reason.strip()[:300],
                            "status": "pending", "applied_at": now_str(), "decided_at": None,
                            "admin_note": ""})
    return days


# ------------------------------------------------------------------ correction requests
def request_correction(emp_id, day, check_in, check_out, reason):
    """Employee asks the admin to fix a day (camera missed them, forgot to mark OUT...)."""
    today = date.today()
    if day > today:
        raise ValueError("A future date can't be corrected.")
    if day < today - timedelta(days=30):
        raise ValueError("Only the last 30 days can be corrected.")
    if not check_in and not check_out:
        raise ValueError("Enter the correct IN and/or OUT time.")
    cin = parse_clock(check_in).strftime("%H:%M") if check_in else ""
    cout = parse_clock(check_out).strftime("%H:%M") if check_out else ""
    if cin and cout and cout <= cin:
        raise ValueError("OUT time must be after IN time.")
    if not reason.strip():
        raise ValueError("Write a reason.")
    if db().corrections.find_one({"emp_id": emp_id, "date": day.isoformat(), "status": "pending"}):
        raise ValueError("You already have a pending correction for this date.")
    db().corrections.insert_one({"emp_id": emp_id, "date": day.isoformat(), "check_in": cin,
                                 "check_out": cout, "reason": reason.strip()[:300], "status": "pending",
                                 "applied_at": now_str(), "decided_at": None, "admin_note": ""})


def apply_correction(req):
    """Admin approved a correction: write it as a manual entry for that day."""
    row = db().attendance.find_one({"emp_id": req["emp_id"], "date": req["date"]}) or {}
    cin = req.get("check_in") or (row.get("check_in") or "")[:5]
    cout = req.get("check_out") or (row.get("check_out") or "")[:5]
    if not cin:
        raise ValueError("No IN time for that day - ask the employee for it, or use Manual entry.")
    manual_entry(req["emp_id"], req["date"], cin, cout, f"Correction: {req['reason']}"[:200])


def today_status(emp_id):
    """IN / last OUT / hours so far for the employee panel."""
    now = datetime.now()
    row = db().attendance.find_one({"emp_id": emp_id, "date": now.strftime("%Y-%m-%d")})
    if row is None:
        rules = Rules()
        if rules.day_off(now.date()):
            status = rules.day_off(now.date())
        elif approved_leaves(now.date(), now.date(), emp_id):
            status = "On Leave"
        else:
            status = "Not In Yet"
        return {"in": "", "out": "", "hours": "", "next": "IN", "status": status, "late_min": 0}
    ev = Rules().evaluate(now.date(), row["check_in"], None, now.date())
    cin = datetime.combine(now.date(), parse_clock(row["check_in"]))
    end = datetime.combine(now.date(), parse_clock(row["check_out"])) if row.get("check_out") else now
    return {"in": row["check_in"][:5], "out": (row.get("check_out") or "")[:5],
            "hours": round(max(0.0, (end - cin).total_seconds() / 3600), 2), "next": "OUT",
            "status": ev["status"], "late_min": ev["late_min"]}


def build_report(start, end, emp_id=None):
    """start/end are date objects. Returns (day rows, per-employee summary)."""
    rules = Rules()
    today = date.today()
    end = min(end, today)
    emps = employees({"_id": emp_id} if emp_id else None)
    att_query = {"date": {"$gte": start.isoformat(), "$lte": end.isoformat()}}
    if emp_id:
        att_query["emp_id"] = emp_id
    recs = {(r["emp_id"], r["date"]): r for r in db().attendance.find(att_query)}
    moves = _movements(start.isoformat(), end.isoformat())
    leaves = approved_leaves(start, end, emp_id)

    rows, summary = [], []
    for emp in emps:
        joined = datetime.strptime(emp["created_at"][:10], "%Y-%m-%d").date()
        sm = {"emp_id": emp["id"], "name": emp["name"], "code": emp["emp_code"],
              "department": emp.get("department", ""), "working_days": 0, "present": 0,
              "on_time": 0, "late": 0, "half_day": 0, "absent": 0, "leave": 0.0,
              "late_min": 0, "hours": 0.0}
        for d in _daterange(start, end):
            ds = d.isoformat()
            r = recs.get((emp["id"], ds))
            row = {"date": ds, "weekday": d.strftime("%a"), "emp_id": emp["id"],
                   "name": emp["name"], "code": emp["emp_code"],
                   "department": emp.get("department", ""), "check_in": "", "check_out": "",
                   "hours": "", "late_min": 0, "in_photo": None, "out_photo": None,
                   "source": "", "note": "", "moves": moves.get((emp["id"], ds), []), "leave": None}
            row["out_count"] = sum(1 for m in row["moves"] if m["kind"] == "OUT")
            lv = leaves.get((emp["id"], ds))
            if lv:
                row["leave"] = lv["type"] + (" (half day)" if lv["half"] else "")
                if r and lv["half"]:
                    sm["leave"] += 0.5
            if r:
                ev = rules.evaluate(d, r["check_in"], r.get("check_out"), today)
                row.update(check_in=_fmt_clock(r["check_in"]), check_out=_fmt_clock(r.get("check_out")),
                           hours=ev["hours"] if ev["hours"] is not None else "",
                           late_min=ev["late_min"], status=ev["status"],
                           in_photo=r.get("in_photo"), out_photo=r.get("out_photo"),
                           source=r.get("source", "camera"), note=r.get("note", ""))
                sm["working_days"] += 1
                sm["present"] += 1
                sm[{"On Time": "on_time", "Late": "late", "Half Day": "half_day"}[ev["status"]]] += 1
                sm["late_min"] += ev["late_min"]
                sm["hours"] += ev["hours"] or 0
            else:
                off = rules.day_off(d)
                if d < joined or not emp.get("active", True):
                    continue
                if off:
                    row["status"] = off
                elif lv:
                    row["status"] = "Leave"
                    sm["working_days"] += 1
                    sm["leave"] += 0.5 if lv["half"] else 1
                    if lv["half"] and d < today:   # the other half of the day was not worked
                        sm["absent"] += 0.5
                elif d == today:
                    row["status"] = "Not In Yet"
                else:
                    row["status"] = "Absent"
                    sm["working_days"] += 1
                    sm["absent"] += 1
            rows.append(row)
        sm["hours"] = round(sm["hours"], 1)
        summary.append(sm)
    rows.sort(key=lambda x: (x["date"], x["name"]), reverse=True)
    return rows, summary


def today_overview():
    rules = Rules()
    now = datetime.now()
    today = now.date()
    ds = today.isoformat()
    everyone = {e["id"]: e for e in employees()}
    emps = [e for e in everyone.values() if e.get("active", True)]
    recs = {r["emp_id"]: r for r in db().attendance.find({"date": ds})}
    moves = _movements(ds, ds)
    on_leave = approved_leaves(today, today)
    events = []
    for e in db().events.find({"ts": {"$gte": ds}}).sort("ts", -1).limit(20):
        who = everyone.get(e["emp_id"])
        if who:
            events.append(dict(e, name=who["name"], emp_code=who["emp_code"]))

    off = rules.day_off(today)
    absent_label = "Absent" if now.time() > rules.half_day_after else "Not In Yet"
    rows = []
    counts = {"total": len(emps), "present": 0, "on_time": 0, "late": 0, "half_day": 0, "not_in": 0,
              "on_leave": 0}
    for emp in emps:
        r = recs.get(emp["id"])
        row = {"emp_id": emp["id"], "name": emp["name"], "code": emp["emp_code"],
               "department": emp.get("department", ""), "check_in": "", "check_out": "",
               "hours": "", "late_min": 0, "in_photo": None, "out_photo": None, "source": "",
               "moves": moves.get((emp["id"], ds), [])}
        if r:
            ev = rules.evaluate(today, r["check_in"], r.get("check_out"), today)
            row.update(check_in=_fmt_clock(r["check_in"]), check_out=_fmt_clock(r.get("check_out")),
                       hours=ev["hours"] if ev["hours"] is not None else "",
                       late_min=ev["late_min"], status=ev["status"], in_photo=r.get("in_photo"),
                       out_photo=r.get("out_photo"), source=r.get("source", "camera"))
            counts["present"] += 1
            counts[{"On Time": "on_time", "Late": "late", "Half Day": "half_day"}[ev["status"]]] += 1
        elif not off and (emp["id"], ds) in on_leave:
            row["status"] = "On Leave"
            counts["on_leave"] += 1
        else:
            row["status"] = off or absent_label
            counts["not_in"] += 1
        rows.append(row)
    rows.sort(key=lambda x: (x["check_in"] == "", x["check_in"], x["name"]))
    return {
        "date": ds,
        "date_label": today.strftime("%A, %d %b %Y"),
        "now": now.strftime("%H:%M:%S"),
        "day_off": off,
        "office_start": rules.office_start.strftime("%H:%M"),
        "grace": rules.grace,
        "late_after": rules.late_after.strftime("%H:%M"),
        "counts": counts,
        "rows": rows,
        "events": [{"name": e["name"], "code": e["emp_code"], "kind": e["kind"],
                    "time": e["ts"][11:16], "photo": e["photo"], "score": e["score"]}
                   for e in events],
    }


def export_excel(rows, summary, start, end, company):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    fills = {"On Time": "DCFCE7", "Late": "FEF3C7", "Half Day": "EDE9FE", "Absent": "FEE2E2",
             "Weekly Off": "F1F5F9", "Holiday": "E0F2FE", "Not In Yet": "F1F5F9", "Leave": "CFFAFE"}
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1E293B")

    def header(ws, cols):
        ws.append(cols)
        for c in ws[1]:
            c.font, c.fill = head_font, head_fill
            c.alignment = Alignment(horizontal="center")
        ws.freeze_panes = "A2"

    def autosize(ws):
        for i, col in enumerate(ws.columns, 1):
            width = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[get_column_letter(i)].width = min(40, max(10, width + 2))

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    header(ws, ["Name", "Emp ID", "Department", "Working Days", "Present", "On Time", "Late",
                "Half Day", "Leave", "Absent", "Total Late (min)", "Total Hours"])
    for s in summary:
        ws.append([s["name"], s["code"], s["department"], s["working_days"], s["present"],
                   s["on_time"], s["late"], s["half_day"], s["leave"], s["absent"], s["late_min"],
                   s["hours"]])
    autosize(ws)

    ws2 = wb.create_sheet("Daily")
    header(ws2, ["Date", "Day", "Name", "Emp ID", "Department", "Check In", "Last Out",
                 "Hours", "Status", "Late (min)", "Times Out", "Movements", "Source", "Note"])
    for r in sorted(rows, key=lambda x: (x["date"], x["name"])):
        ws2.append([r["date"], r["weekday"], r["name"], r["code"], r["department"], r["check_in"],
                    r["check_out"], r["hours"], r["status"], r["late_min"] or "",
                    r["out_count"] or "", moves_text(r["moves"]), r["source"],
                    "; ".join(x for x in (r.get("leave") or "", r["note"]) if x)])
        fill = fills.get(r["status"])
        if fill:
            ws2.cell(row=ws2.max_row, column=9).fill = PatternFill("solid", fgColor=fill)
    autosize(ws2)

    ws.insert_rows(1)
    ws["A1"] = f"{company} - Attendance {start} to {end}"
    ws["A1"].font = Font(bold=True, size=13)
    ws.freeze_panes = "A3"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
